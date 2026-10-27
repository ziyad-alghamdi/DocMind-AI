from io import BytesIO
import os
import re
from pathlib import Path
from typing import Annotated
from uuid import uuid4

import httpx
import numpy as np
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_BASE_URL = "https://generativelanguage.googleapis.com/v1beta"
EMBEDDING_MODEL = "gemini-embedding-2"
ANSWER_MODEL = "gemini-3.7-flash"
EMBEDDING_DIMENSIONS = 768
TIMEOUT = 300.0

allowed_origins_env = os.getenv("ALLOWED_ORIGINS", "")
ALLOWED_ORIGINS = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
]
if allowed_origins_env:
    ALLOWED_ORIGINS.extend(
        origin.strip()
        for origin in allowed_origins_env.split(",")
        if origin.strip()
    )

app = FastAPI(
    title="DocMind AI API",
    description="Cloud-ready multilingual RAG powered by Gemini",
    version="4.0.0",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DOCUMENTS = {}
CHUNKS = []

class AskRequest(BaseModel):
    question: str = Field(min_length=2)
    document_ids: list[str] | None = None
    top_k: int = Field(default=5, ge=1, le=10)

class SearchRequest(BaseModel):
    question: str = Field(min_length=2)
    document_ids: list[str] | None = None
    top_k: int = Field(default=5, ge=1, le=10)

class SummaryRequest(BaseModel):
    document_ids: list[str] | None = None
    language: str = "auto"


def gemini_headers():
    if not GEMINI_API_KEY:
        raise HTTPException(
            status_code=503,
            detail="Gemini API key is not configured on the server.",
        )
    return {
        "Content-Type": "application/json",
        "x-goog-api-key": GEMINI_API_KEY,
    }


def check_gemini():
    return bool(GEMINI_API_KEY)


def create_embeddings(texts: list[str], task: str = "document"):
    if not texts:
        return []

    if task == "query":
        prepared = [f"task: retrieval_query | query: {item}" for item in texts]
    else:
        prepared = [f"task: retrieval_document | document: {item}" for item in texts]

    result = []
    try:
        with httpx.Client(timeout=TIMEOUT) as client:
            for i in range(0, len(prepared), 20):
                batch = prepared[i:i + 20]
                payload = {
                    "requests": [
                        {
                            "model": f"models/{EMBEDDING_MODEL}",
                            "content": {"parts": [{"text": item}]},
                            "outputDimensionality": EMBEDDING_DIMENSIONS,
                        }
                        for item in batch
                    ]
                }
                response = client.post(
                    f"{GEMINI_BASE_URL}/models/{EMBEDDING_MODEL}:batchEmbedContents",
                    headers=gemini_headers(),
                    json=payload,
                )
                response.raise_for_status()
                embeddings = response.json().get("embeddings", [])
                if len(embeddings) != len(batch):
                    raise RuntimeError("Unexpected embeddings response from Gemini")
                result.extend(item["values"] for item in embeddings)
        return result
    except httpx.HTTPStatusError as error:
        print("Gemini embedding HTTP error:", error.response.text)
        raise HTTPException(
            status_code=502,
            detail="Gemini embedding request failed. Check the API key and free-tier quota.",
        )
    except HTTPException:
        raise
    except Exception as error:
        print("Embedding error:", repr(error))
        raise HTTPException(502, "Could not create document embeddings.")


def gemini_chat(system_prompt: str, user_prompt: str):
    full_prompt = f"{system_prompt}\n\n---\n\n{user_prompt}"
    try:
        with httpx.Client(timeout=TIMEOUT) as client:
            response = client.post(
                f"{GEMINI_BASE_URL}/models/{ANSWER_MODEL}:generateContent",
                headers=gemini_headers(),
                json={
                    "contents": [
                        {
                            "role": "user",
                            "parts": [{"text": full_prompt}],
                        }
                    ],
                    "generationConfig": {
                        "temperature": 0.0,
                        "maxOutputTokens": 2048,
                    },
                },
            )
            response.raise_for_status()
            data = response.json()
            candidates = data.get("candidates", [])
            if not candidates:
                raise RuntimeError("Gemini returned no candidates")
            parts = candidates[0].get("content", {}).get("parts", [])
            answer = "".join(part.get("text", "") for part in parts).strip()
            if not answer:
                raise RuntimeError("Gemini returned an empty answer")
            return answer
    except httpx.HTTPStatusError as error:
        print("Gemini chat HTTP error:", error.response.text)
        raise HTTPException(
            status_code=502,
            detail="Gemini answer request failed. Check the API key and free-tier quota.",
        )
    except HTTPException:
        raise
    except Exception as error:
        print("Chat error:", repr(error))
        raise HTTPException(502, "Could not generate an AI answer.")


def cosine_similarity(a, b):
    a = np.array(a, dtype=np.float32)
    b = np.array(b, dtype=np.float32)
    den = np.linalg.norm(a) * np.linalg.norm(b)
    return 0.0 if den == 0 else float(np.dot(a, b) / den)


def extract_pdf(data: bytes):
    reader = PdfReader(BytesIO(data))
    out = []
    for n, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append({"location": f"Page {n}", "location_number": n, "text": text})
    return out


def extract_docx(data: bytes):
    doc = Document(BytesIO(data))
    out = []
    n = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            n += 1
            out.append({"location": f"Paragraph {n}", "location_number": n, "text": text})
    return out


def extract_pptx(data: bytes):
    prs = Presentation(BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    parts.append(text)
        if parts:
            out.append({"location": f"Slide {n}", "location_number": n, "text": "\n".join(parts)})
    return out


def chunk_text(text: str, max_chars=1400, overlap=250):
    text = text.strip()
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]

    out = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            search_start = start + int(max_chars * 0.60)
            section = text[search_start:end]
            breaks = [
                section.rfind("\n"), section.rfind(". "), section.rfind("? "),
                section.rfind("! "), section.rfind("، "), section.rfind("؟ ")
            ]
            best = max(breaks)
            if best > 0:
                end = search_start + best + 1

        piece = text[start:end].strip()
        if piece:
            out.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return out


def create_document_chunks(document_id: str, filename: str, locations: list):
    out = []
    for loc in locations:
        for piece_number, piece in enumerate(chunk_text(loc["text"]), 1):
            out.append({
                "id": str(uuid4()),
                "document_id": document_id,
                "filename": filename,
                "location": loc["location"],
                "location_number": loc["location_number"],
                "piece_number": piece_number,
                "text": piece,
                "embedding": None,
            })
    return out



FIELD_INTENTS = [
    {
        "question_terms": ["الاسم", "اسم", "name", "full name"],
        "document_terms": ["personal information", "full name", "name"],
    },
    {
        "question_terms": ["الجامعة", "جامعة", "university", "college"],
        "document_terms": ["university", "college", "education"],
    },
    {
        "question_terms": ["الجوال", "الهاتف", "رقم الجوال", "phone", "mobile", "telephone"],
        "document_terms": ["phone number", "phone", "mobile", "telephone"],
    },
    {
        "question_terms": ["البريد", "الايميل", "الإيميل", "email", "e-mail"],
        "document_terms": ["e-mail", "email"],
    },
    {
        "question_terms": ["الجنسية", "nationality"],
        "document_terms": ["nationality"],
    },
    {
        "question_terms": ["الموقع", "المدينة", "location", "city"],
        "document_terms": ["location", "city"],
    },
    {
        "question_terms": ["التخصص", "major", "specialization"],
        "document_terms": ["major", "computer science", "specialization", "faculty"],
    },
    {
        "question_terms": ["الخبرات", "الخبرة", "experience", "work experience"],
        "document_terms": ["experience", "work", "internship", "training"],
    },
    {
        "question_terms": ["المهارات", "skills", "skill"],
        "document_terms": ["skills", "skill"],
    },
]


def field_intent_boost(question: str, text: str) -> float:
    """Small lexical boost for explicit factual fields in Arabic/English."""
    q = question.lower()
    t = text.lower()

    best_boost = 0.0

    for intent in FIELD_INTENTS:
        if any(term in q for term in intent["question_terms"]):
            if any(term in t for term in intent["document_terms"]):
                best_boost = max(best_boost, 0.22)

    return best_boost


def semantic_search(question: str, document_ids=None, top_k=5):
    if not CHUNKS:
        raise HTTPException(400, "No processed documents found. Upload documents first.")

    candidates = CHUNKS
    if document_ids:
        candidates = [c for c in CHUNKS if c["document_id"] in document_ids]
    if not candidates:
        raise HTTPException(404, "No chunks found for the selected documents.")

    q = create_embeddings([question], task="query")[0]
    results = []
    for chunk in candidates:
        similarity = cosine_similarity(q, chunk["embedding"])
        lexical_boost = field_intent_boost(question, chunk["text"])
        results.append({
            **chunk,
            "similarity": similarity,
            "hybrid_score": similarity + lexical_boost,
        })
    results.sort(key=lambda x: x["hybrid_score"], reverse=True)
    return results[:top_k]


@app.get("/")
def root():
    return {
        "message": "DocMind AI API is running",
        "version": "4.0.0",
        "ai_provider": "Gemini API",
        "embedding_model": EMBEDDING_MODEL,
        "answer_model": ANSWER_MODEL,
        "documents": len(DOCUMENTS),
        "chunks": len(CHUNKS),
    }


@app.get("/health")
def health():
    return {
        "status": "healthy",
        "gemini_configured": check_gemini(),
        "embedding_model": EMBEDDING_MODEL,
        "answer_model": ANSWER_MODEL,
        "documents": len(DOCUMENTS),
        "chunks": len(CHUNKS),
    }


@app.post("/upload")
async def upload_documents(
    files: Annotated[list[UploadFile], File(description="Upload PDF, DOCX, or PPTX files")]
):
    uploaded = []

    for file in files:
        if not file.filename:
            continue

        ext = Path(file.filename).suffix.lower()
        if ext not in [".pdf", ".docx", ".pptx"]:
            raise HTTPException(400, f"Unsupported file type: {ext}")

        data = await file.read()
        try:
            if ext == ".pdf":
                locations = extract_pdf(data)
            elif ext == ".docx":
                locations = extract_docx(data)
            else:
                locations = extract_pptx(data)
        except Exception as e:
            print("Extraction error:", file.filename, repr(e))
            raise HTTPException(400, f"Could not read {file.filename}")

        if not locations:
            raise HTTPException(
                400,
                f"No readable text was found in {file.filename}. Scanned PDFs need OCR support.",
            )

        document_id = str(uuid4())
        document_chunks = create_document_chunks(document_id, file.filename, locations)
        embeddings = create_embeddings([c["text"] for c in document_chunks])

        for chunk, embedding in zip(document_chunks, embeddings):
            chunk["embedding"] = embedding

        DOCUMENTS[document_id] = {
            "id": document_id,
            "filename": file.filename,
            "file_type": ext.replace(".", ""),
            "locations": locations,
        }
        CHUNKS.extend(document_chunks)

        uploaded.append({
            "id": document_id,
            "filename": file.filename,
            "file_type": ext.replace(".", ""),
            "locations_found": len(locations),
            "characters_extracted": sum(len(x["text"]) for x in locations),
            "chunks_created": len(document_chunks),
            "preview": locations[0]["text"][:250],
        })

    return {"success": True, "documents": uploaded, "total_chunks": len(CHUNKS)}


@app.get("/documents")
def get_documents():
    return [
        {
            "id": d["id"],
            "filename": d["filename"],
            "file_type": d["file_type"],
            "locations_found": len(d["locations"]),
        }
        for d in DOCUMENTS.values()
    ]


@app.get("/documents/{document_id}")
def get_document(document_id: str):
    doc = DOCUMENTS.get(document_id)
    if not doc:
        raise HTTPException(404, "Document not found")
    return doc


@app.post("/search")
def search_documents(request: SearchRequest):
    results = semantic_search(request.question, request.document_ids, request.top_k)
    return {
        "question": request.question,
        "results": [
            {
                "document_id": r["document_id"],
                "filename": r["filename"],
                "location": r["location"],
                "similarity": round(r["similarity"], 4),
                "text": r["text"],
            }
            for r in results
        ],
    }


@app.post("/ask")
def ask_documents(request: AskRequest):
    results = semantic_search(request.question, request.document_ids, request.top_k)

    context_parts = []
    sources = []
    seen = set()

    for i, r in enumerate(results, 1):
        context_parts.append(
            f'[SOURCE {i}]\nFile: {r["filename"]}\nLocation: {r["location"]}\n\n{r["text"]}'
        )
        key = (r["filename"], r["location"])
        if key not in seen:
            seen.add(key)
            sources.append({
                "source_number": i,
                "document_id": r["document_id"],
                "filename": r["filename"],
                "location": r["location"],
                "similarity": round(r["similarity"], 4),
                "preview": r["text"][:300],
            })

    system_prompt = """
You are DocMind AI, a strict document-grounded assistant.

NON-NEGOTIABLE RULES:
1. Use ONLY facts explicitly written in the retrieved document context.
2. Do NOT infer, guess, complete, or reinterpret missing facts.
3. Before answering, locate direct textual evidence that answers the exact question.
4. For exact factual fields such as a person's name, phone number, email, university,
   nationality, date, ID, or location, COPY the exact value from the source text.
   Do not alter names, numbers, spellings, initials, or email addresses.
5. A project title or project description is NOT evidence for an unrelated personal fact.
6. If no retrieved source directly states the requested fact, say that you could not
   find this information in the uploaded documents.
7. Answer in the same language as the user's question unless the user asks otherwise.
8. The source may be in another language; translation is allowed, but exact names,
   numbers, emails, and identifiers must stay exactly as written in the source.
9. Cite every factual answer with [1], [2], etc.
10. Cite ONLY sources that directly support the answer.
11. Be concise for simple factual questions.

Example behavior:
Question: "ما هو الاسم؟"
If a source explicitly shows "Ziyad Mohammed A Alghamdi", answer:
"الاسم هو Ziyad Mohammed A Alghamdi. [1]"
Do not invent an Arabic surname or change the spelling.
""".strip()

    user_prompt = f"""
USER QUESTION:
{request.question}

RETRIEVED DOCUMENT CONTEXT:
{chr(10).join(context_parts)}

Answer using only the context above.
""".strip()

    answer = gemini_chat(system_prompt, user_prompt)

    # Return only sources actually cited by the model.
    cited_numbers = {
        int(number)
        for number in re.findall(r"\[(\d+)\]", answer)
    }

    cited_sources = [
        source
        for source in sources
        if source["source_number"] in cited_numbers
    ]

    return {
        "question": request.question,
        "answer": answer,
        "sources": cited_sources,
    }


@app.post("/summarize")
def summarize_documents(request: SummaryRequest):
    selected = list(DOCUMENTS.values())
    if request.document_ids:
        selected = [d for d in selected if d["id"] in request.document_ids]
    if not selected:
        raise HTTPException(404, "No documents found to summarize.")

    parts = []
    total = 0
    max_chars = 80000
    for d in selected:
        parts.append(f'\n=== FILE: {d["filename"]} ===\n')
        for loc in d["locations"]:
            section = f'\n{loc["location"]}:\n{loc["text"]}\n'
            if total + len(section) > max_chars:
                break
            parts.append(section)
            total += len(section)
        if total >= max_chars:
            break

    lang = request.language.lower()
    if lang == "arabic":
        lang_instruction = "Write the summary in Arabic."
    elif lang == "english":
        lang_instruction = "Write the summary in English."
    else:
        lang_instruction = "Write the summary in the main language of the documents."

    system_prompt = f"""
You are DocMind AI.
Summarize ONLY the supplied documents. Do not add outside knowledge.
Create a clear structured summary with main ideas, important concepts, key facts,
important definitions when present, and important conclusions when present.
{lang_instruction}
""".strip()

    summary = gemini_chat(system_prompt, "".join(parts))
    return {
        "summary": summary,
        "documents": [{"id": d["id"], "filename": d["filename"]} for d in selected],
    }
